"""
Audit log for enrollment operations — PostgreSQL via db.py.
"""
from __future__ import annotations

import io
import uuid
from datetime import datetime, timezone
from typing import Any

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

import db

_CREATE_AUDIT = """
CREATE TABLE IF NOT EXISTS audit_log (
    id           SERIAL PRIMARY KEY,
    timestamp    TEXT    NOT NULL,
    ejecucion_id TEXT    NOT NULL,
    cedula       TEXT    NOT NULL,
    nombre       TEXT    NOT NULL,
    accion       TEXT    NOT NULL,
    plataforma   TEXT,
    curso        TEXT,
    semestre     TEXT,
    detalle      TEXT
)
"""

_CREATE_EJECUCIONES = """
CREATE TABLE IF NOT EXISTS ejecuciones (
    id            TEXT PRIMARY KEY,
    timestamp     TEXT NOT NULL,
    tipo          TEXT,
    semestre      TEXT,
    total         INTEGER DEFAULT 0,
    creados       INTEGER DEFAULT 0,
    existentes    INTEGER DEFAULT 0,
    inscripciones INTEGER DEFAULT 0,
    errores       INTEGER DEFAULT 0,
    duracion_seg  REAL,
    estado        TEXT
)
"""


async def init_db() -> None:
    import user_service
    import webhook_service
    import course_matcher
    await db.execute(_CREATE_AUDIT)
    await db.execute(_CREATE_EJECUCIONES)
    await user_service.init_users_table()
    await webhook_service.init_pending_table()
    await course_matcher.init_matcher_tables()


async def log_action(
    ejecucion_id: str,
    cedula: str,
    nombre: str,
    accion: str,
    plataforma: str = "",
    curso: str = "",
    semestre: str = "",
    detalle: str = "",
) -> None:
    ts = datetime.now(timezone.utc).isoformat()
    await db.execute(
        """INSERT INTO audit_log
           (timestamp, ejecucion_id, cedula, nombre, accion, plataforma, curso, semestre, detalle)
           VALUES (?,?,?,?,?,?,?,?,?)""",
        ts, ejecucion_id, cedula, nombre, accion, plataforma, curso, semestre, detalle,
    )


async def save_ejecucion(summary: dict) -> None:
    await db.execute(
        """INSERT INTO ejecuciones
           (id, timestamp, tipo, semestre, total, creados, existentes, inscripciones, errores, duracion_seg, estado)
           VALUES (?,?,?,?,?,?,?,?,?,?,?)
           ON CONFLICT (id) DO UPDATE SET
             timestamp=EXCLUDED.timestamp, tipo=EXCLUDED.tipo, semestre=EXCLUDED.semestre,
             total=EXCLUDED.total, creados=EXCLUDED.creados, existentes=EXCLUDED.existentes,
             inscripciones=EXCLUDED.inscripciones, errores=EXCLUDED.errores,
             duracion_seg=EXCLUDED.duracion_seg, estado=EXCLUDED.estado""",
        summary.get("ejecucion_id", str(uuid.uuid4())),
        summary.get("timestamp", datetime.now(timezone.utc).isoformat()),
        summary.get("tipo", "manual"),
        summary.get("semestre", ""),
        summary.get("total", 0),
        summary.get("creados", 0),
        summary.get("existentes", 0),
        summary.get("inscripciones", 0),
        summary.get("errores", 0),
        summary.get("duracion_seg", 0.0),
        summary.get("estado", "exitoso"),
    )


async def get_historial(
    semestre: str | None = None,
    cedula: str | None = None,
    limit: int = 500,
) -> list[dict]:
    conditions: list[str] = []
    params: list[Any] = []
    if semestre:
        conditions.append("semestre = ?")
        params.append(semestre)
    if cedula:
        conditions.append("cedula = ?")
        params.append(cedula)
    where = ("WHERE " + " AND ".join(conditions)) if conditions else ""
    params.append(limit)
    return await db.fetch(
        f"SELECT * FROM audit_log {where} ORDER BY timestamp DESC LIMIT ?", *params
    )


async def get_ejecuciones(limit: int = 50) -> list[dict]:
    return await db.fetch(
        "SELECT * FROM ejecuciones ORDER BY timestamp DESC LIMIT ?", limit
    )


async def get_dashboard_kpis() -> dict:
    today = datetime.now(timezone.utc).date().isoformat()

    rows = await db.fetch(
        "SELECT accion, COUNT(*) as cnt FROM audit_log WHERE LEFT(timestamp, 10) = ? GROUP BY accion",
        today,
    )
    today_by_action: dict[str, int] = {r["accion"]: r["cnt"] for r in rows}

    daily = await db.fetch(
        """SELECT LEFT(timestamp, 10) as day, COUNT(DISTINCT cedula) as cnt
           FROM audit_log WHERE LEFT(timestamp, 10) >= ?
           GROUP BY day ORDER BY day""",
        (datetime.now(timezone.utc).date() - __import__('datetime').timedelta(days=6)).isoformat(),
    )

    ejecuciones = await db.fetch(
        "SELECT * FROM ejecuciones ORDER BY timestamp DESC LIMIT 10"
    )

    total_alumnos = await db.fetchval(
        "SELECT COUNT(DISTINCT cedula) FROM audit_log"
    ) or 0

    return {
        "hoy": {
            "procesados": sum(today_by_action.values()),
            "creados": today_by_action.get("creado", 0),
            "inscritos": today_by_action.get("inscrito", 0),
            "errores": today_by_action.get("error", 0),
        },
        "total_alumnos": total_alumnos,
        "ultimos_7_dias": [{"day": r["day"], "count": r["cnt"]} for r in daily],
        "ultimas_ejecuciones": ejecuciones,
    }


def _header_row(ws, headers: list[str], color: str = "1E3A5F") -> None:
    fill = PatternFill("solid", fgColor=color)
    font = Font(bold=True, color="FFFFFF")
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)


def _autowidth(ws, max_width: int = 55) -> None:
    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 3, max_width)


async def export_to_excel(semestre: str | None = None, cedula: str | None = None) -> bytes:
    detail_rows = await get_historial(semestre=semestre, cedula=cedula, limit=10000)
    ejecuciones = await get_ejecuciones(limit=200)

    wb = openpyxl.Workbook()
    ws_exec = wb.active
    ws_exec.title = "Ejecuciones"

    exec_headers = [
        "Fecha", "Hora (UTC)", "ID Ejecución", "Tipo", "Semestre",
        "Total alumnos", "Creados", "Existentes", "Inscripciones",
        "Errores", "Duración (seg)", "Estado",
    ]
    _header_row(ws_exec, exec_headers, "1E3A5F")

    estado_fills = {
        "exitoso":          PatternFill("solid", fgColor="D1FAE5"),
        "con_errores":      PatternFill("solid", fgColor="FEF3C7"),
        "error_validacion": PatternFill("solid", fgColor="FEE2E2"),
        "fallido":          PatternFill("solid", fgColor="FEE2E2"),
        "sin_datos":        PatternFill("solid", fgColor="F3F4F6"),
    }

    for r_idx, ej in enumerate(ejecuciones, 2):
        ts = ej.get("timestamp", "")
        try:
            dt = datetime.fromisoformat(ts.replace("Z", "+00:00"))
            fecha = dt.strftime("%d/%m/%Y")
            hora = dt.strftime("%H:%M:%S")
        except Exception:
            fecha, hora = ts, ""
        estado = ej.get("estado", "")
        fill = estado_fills.get(estado, PatternFill("solid", fgColor="FFFFFF"))
        values = [
            fecha, hora, ej.get("id", ""), ej.get("tipo", ""), ej.get("semestre", ""),
            ej.get("total", 0), ej.get("creados", 0), ej.get("existentes", 0),
            ej.get("inscripciones", 0), ej.get("errores", 0), ej.get("duracion_seg", 0), estado,
        ]
        for c_idx, val in enumerate(values, 1):
            cell = ws_exec.cell(row=r_idx, column=c_idx, value=val)
            cell.fill = fill
            if c_idx == 12:
                font_color = "065F46" if estado == "exitoso" else "991B1B" if estado in ("fallido", "error_validacion") else "92400E"
                cell.font = Font(bold=True, color=font_color)
    _autowidth(ws_exec)

    ws_det = wb.create_sheet("Detalle")
    det_headers = [
        "Fecha", "Hora (UTC)", "Cédula", "Nombre", "Acción",
        "Plataforma", "Curso / Materia", "Semestre", "Detalle / Error", "ID Ejecución",
    ]
    _header_row(ws_det, det_headers, "374151")
    action_fills = {
        "creado":   PatternFill("solid", fgColor="D1FAE5"),
        "inscrito": PatternFill("solid", fgColor="DBEAFE"),
        "error":    PatternFill("solid", fgColor="FEE2E2"),
        "omitido":  PatternFill("solid", fgColor="F3F4F6"),
    }
    for r_idx, row in enumerate(detail_rows, 2):
        ts = row.get("timestamp", "")
        try:
            dt = datetime.fromisoformat(ts.replace("Z", "+00:00"))
            fecha = dt.strftime("%d/%m/%Y")
            hora = dt.strftime("%H:%M:%S")
        except Exception:
            fecha, hora = ts, ""
        accion = row.get("accion", "")
        fill = action_fills.get(accion, PatternFill("solid", fgColor="FFFFFF"))
        detalle = row.get("detalle", "")
        values = [
            fecha, hora, row.get("cedula", ""), row.get("nombre", ""), accion,
            row.get("plataforma", ""), row.get("curso", ""), row.get("semestre", ""),
            detalle, row.get("ejecucion_id", ""),
        ]
        for c_idx, val in enumerate(values, 1):
            cell = ws_det.cell(row=r_idx, column=c_idx, value=val)
            cell.fill = fill
            if c_idx == 5 and accion == "error":
                cell.font = Font(bold=True, color="991B1B")
            if c_idx == 9 and detalle:
                cell.alignment = Alignment(wrap_text=True)
    _autowidth(ws_det)
    ws_det.row_dimensions[1].height = 30

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
