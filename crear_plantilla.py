"""Script de un solo uso para generar plantilla_carga_masiva.xlsx"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

wb = Workbook()

# ---- Hoja Usuarios ----
ws_u = wb.active
ws_u.title = "Usuarios"

headers_u = ["nombre", "email", "sis_id", "rol_canvas", "upn_azure", "grupo_azure", "equipo_teams"]
examples_u = [
    ["Juan Pérez", "juan.perez@uni.edu", "STU-001", "StudentEnrollment",
     "juan.perez@uni.onmicrosoft.com", "grupo-id-azure", "equipo-id-teams"],
    ["María López", "maria.lopez@uni.edu", "STU-002", "StudentEnrollment",
     "maria.lopez@uni.onmicrosoft.com", "", ""],
]

header_fill = PatternFill("solid", fgColor="1F4E79")
example_fill = PatternFill("solid", fgColor="D9E1F2")

for col, h in enumerate(headers_u, 1):
    cell = ws_u.cell(row=1, column=col, value=h)
    cell.fill = header_fill
    cell.font = Font(color="FFFFFF", bold=True)
    cell.alignment = Alignment(horizontal="center")

for row_idx, row in enumerate(examples_u, 2):
    for col, val in enumerate(row, 1):
        cell = ws_u.cell(row=row_idx, column=col, value=val)
        cell.fill = example_fill

for col in range(1, len(headers_u) + 1):
    ws_u.column_dimensions[get_column_letter(col)].width = 30

# ---- Hoja Inscripciones ----
ws_e = wb.create_sheet("Inscripciones")

headers_e = ["email_usuario", "curso_canvas_id", "rol_canvas", "grupo_azure", "equipo_teams", "canal_teams"]
examples_e = [
    ["juan.perez@uni.edu", "12345", "StudentEnrollment", "grupo-id-azure", "equipo-id-teams", "General"],
    ["maria.lopez@uni.edu", "12345", "StudentEnrollment", "", "", ""],
]

for col, h in enumerate(headers_e, 1):
    cell = ws_e.cell(row=1, column=col, value=h)
    cell.fill = header_fill
    cell.font = Font(color="FFFFFF", bold=True)
    cell.alignment = Alignment(horizontal="center")

for row_idx, row in enumerate(examples_e, 2):
    for col, val in enumerate(row, 1):
        cell = ws_e.cell(row=row_idx, column=col, value=val)
        cell.fill = example_fill

for col in range(1, len(headers_e) + 1):
    ws_e.column_dimensions[get_column_letter(col)].width = 30

wb.save("plantilla_carga_masiva.xlsx")
print("Plantilla generada: plantilla_carga_masiva.xlsx")
